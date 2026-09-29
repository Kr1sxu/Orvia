"""将一个已保存的 M15 回答编排为固定版式简报，并在写入前读回验证。"""

from io import BytesIO
from contextlib import closing
import hashlib
import json
from pathlib import Path

from ..chat.synthesis import Generated
from ..computer.paths import ToolError


LABELS = {"fact": "证据陈述", "inference": "推断", "conflict": "来源冲突", "unknown": "无法回答"}
FONT_PATH = Path(__file__).parent / "assets" / "NotoSansSC.ttf"


def _parts(text: str, size: int) -> list[str]:
    """纯文本按码点分段，不省略原文；页/幻灯片边界在预览与文件中相同。"""
    parts, current, breaks = [], "", 0
    for char in text:
        if len(current) >= size or breaks >= 6:
            parts.append(current)
            current, breaks = "", 0
        current += char
        breaks += char == "\n"
    if current:
        parts.append(current)
    return parts


def _digest(value: dict) -> str:
    return hashlib.sha256(json.dumps(value, ensure_ascii=False, sort_keys=True).encode("utf-8")).hexdigest()


def prepare_publication(message: dict, request) -> dict:
    """只允许覆写标题与纯文字；引用、类型及版本一律来自同一条历史消息。"""
    if message.get("kind") != "synthesis" or not isinstance(message.get("data"), dict):
        raise ToolError("INVALID_SOURCE", "仅能从当前会话已保存的模型综合回答制作简报")
    data = message["data"]
    try:
        original = Generated.model_validate({"answer": data["answer"], "claims": data["claims"]})
        citations = data["citations"]
        coverage = data["coverage"]
    except (KeyError, TypeError, ValueError):
        raise ToolError("INVALID_SOURCE", "已保存回答结构无效，请重新生成") from None
    if len(request.claim_texts) != len(original.claims) or any(not value.strip() or len(value) > 600 for value in request.claim_texts):
        raise ToolError("INVALID_PARAMS", "结论文字数量或长度与原结果不符")
    if not isinstance(citations, list) or not isinstance(coverage, list) or not all(isinstance(item, dict) for item in [*citations, *coverage]):
        raise ToolError("INVALID_SOURCE", "来源引用结构无效")
    if not request.title.strip() or not request.answer.strip():
        raise ToolError("INVALID_PARAMS", "成品标题和摘要正文不能为空")
    edited = [request.title, request.answer, *request.claim_texts]
    if any(0xD800 <= ord(char) <= 0xDFFF or (ord(char) < 32 and char not in "\t\n\r") for value in edited for char in value):
        raise ToolError("INVALID_CONTENT", "成品文字含不支持的控制字符")
    allowed = {item["citation"] for item in citations if isinstance(item, dict) and isinstance(item.get("citation"), str)}
    if len(allowed) != len(citations) or any(cite not in allowed for claim in original.claims for cite in claim.citations):
        raise ToolError("INVALID_SOURCE", "引用映射不完整，请重新生成")
    ref_numbers = {item["citation"]: index + 1 for index, item in enumerate(citations)}
    def printable(value):
        return "".join(char if (ord(char) >= 32 and not 0xD800 <= ord(char) <= 0xDFFF) or char in "\t\n\r" else " " for char in str(value))[:100]
    try:
        titles = {(item["kind"], item["evidence_id"]): printable(item["title"]) for item in coverage}
    except KeyError:
        raise ToolError("INVALID_SOURCE", "来源覆盖结构无效") from None
    references = []
    for item in citations:
        references.append({"number": ref_numbers[item["citation"]], "citation": item["citation"],
                           "title": titles.get((item["kind"], item["evidence_id"]), "来源"),
                           "locator": printable(item["locator"]), "kind": item["kind"], "evidence_id": item["evidence_id"]})
    claims = [{"kind": claim.kind, "label": LABELS[claim.kind], "text": text.strip(),
               "numbers": [ref_numbers[cite] for cite in claim.citations]}
              for claim, text in zip(original.claims, request.claim_texts, strict=True)]
    is_slide = request.format == "pptx"
    size = 180 if is_slide else 500
    pages = []
    for index, fragment in enumerate(_parts(request.answer.strip(), size)):
        pages.append({"heading": "摘要" if index == 0 else f"摘要（续 {index + 1}）", "body": fragment,
                      "label": "模型回答（用户可编辑，需核对引用）", "numbers": []})
    for index, claim in enumerate(claims, 1):
        parts = _parts(claim["text"], size)
        for part_index, fragment in enumerate(parts):
            pages.append({"heading": f"结论 {index}" + (f"（续 {part_index + 1}）" if part_index else ""),
                          "body": fragment, "label": claim["label"], "numbers": claim["numbers"]})
    if not is_slide:
        # 报告把较短的摘要与结论排在同一页，避免一个短句占一整页；每页仍有硬预算。
        grouped, sections, length = [], [], 0
        for section in pages:
            content = f"{section['heading']} · {section['label']}"
            if section["numbers"]:
                content += " " + " ".join(f"[{n}]" for n in section["numbers"])
            content += "\n" + section["body"]
            if sections and length + len(content) > 600:
                grouped.append(sections)
                sections, length = [], 0
            sections.append(content)
            length += len(content)
        if sections:
            grouped.append(sections)
        pages = [{"heading": "正文" if index == 0 else f"正文（续 {index + 1}）",
                  "body": "\n\n".join(parts), "label": "摘要与结论；请按编号回查来源", "numbers": []}
                 for index, parts in enumerate(grouped)]
    # 附录每页最多三个引用；完整 ID 保留，显示编号仅用于正文排版。
    for start in range(0, len(references), 3):
        pages.append({"heading": "来源与引用" + (f"（续 {start // 3 + 1}）" if start else ""),
                      "body": "", "label": "引用仅校验身份和定位，事实仍需人工复核。",
                      "numbers": [], "references": references[start:start + 3]})
    if len(pages) > 48:
        raise ToolError("LAYOUT_LIMIT", "简报超过 48 个版面，请缩短文字或减少手工换行")
    packet = {"schema": "orvia.publication.v1", "format": request.format, "message_id": str(request.message_id),
              "title": request.title.strip(), "answer": request.answer.strip(), "claims": claims,
              "references": references, "pages": pages,
              "notice": "根据当前会话 M15 结果制作；文字可经用户编辑，引用仅映射原证据，未经语义核实。OCR/截断/冲突请回查原文。",
              "source_revision": _digest(data)}
    packet["revision"] = _digest(packet)
    packet["filename"] = f"Orvia-brief-{packet['revision'][:12]}.{request.format}"
    return packet


def _render_docx(packet: dict) -> bytes:
    from docx import Document
    from docx.shared import Cm, Pt

    document = Document()
    section = document.sections[0]
    section.top_margin = section.bottom_margin = Cm(2.2)
    section.left_margin = section.right_margin = Cm(2.4)
    styles = document.styles
    styles["Normal"].font.name = "Microsoft YaHei"
    styles["Normal"].font.size = Pt(10.5)
    document.core_properties.title = packet["title"]
    for index, page in enumerate(packet["pages"]):
        if index:
            document.add_page_break()
        document.add_heading(packet["title"], 0 if index == 0 else 2)
        document.add_heading(page["heading"], 1 if index == 0 else 2)
        document.add_paragraph(page["label"])
        if page["body"]:
            document.add_paragraph(page["body"])
        if page["numbers"]:
            document.add_paragraph("引用：" + " ".join(f"[{n}]" for n in page["numbers"]))
        if page.get("references"):
            table = document.add_table(rows=1, cols=2)
            table.style = "Table Grid"
            table.rows[0].cells[0].text, table.rows[0].cells[1].text = "编号", "来源、定位与证据 ID"
            for ref in page["references"]:
                row = table.add_row().cells
                row[0].text = f"[{ref['number']}]"
                row[1].text = f"{ref['title']} · {ref['locator']}\n{ref['citation']}"
        document.add_paragraph(packet["notice"])
    output = BytesIO()
    document.save(output)
    # 实际用相同库读回，检验包结构与可编辑文字，不以扩展名冒充 Office 成品。
    reopened = Document(BytesIO(output.getvalue()))
    if packet["title"] not in "\n".join(p.text for p in reopened.paragraphs):
        raise ToolError("EXPORT_FAILED", "Word 成品读回校验失败")
    return output.getvalue()


def _render_pptx(packet: dict) -> bytes:
    from pptx import Presentation
    from pptx.util import Inches, Pt

    presentation = Presentation()
    presentation.slide_width, presentation.slide_height = Inches(13.333), Inches(7.5)
    for page in packet["pages"]:
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        for top, height, content, size, bold in (
            (0.45, 0.6, packet["title"], 20, True),
            (1.18, 0.65, page["heading"], 28, True),
            (2.0, 0.38, page["label"], 13, False),
            (2.55, 3.55, page["body"] or "\n".join(f"[{r['number']}] {r['title']} · {r['locator']}\n{r['citation']}" for r in page.get("references", [])), 19 if page["body"] else 13, False),
            (6.33, 0.42, "引用：" + " ".join(f"[{n}]" for n in page["numbers"]) if page["numbers"] else packet["notice"], 10, False),
        ):
            box = slide.shapes.add_textbox(Inches(0.72), Inches(top), Inches(11.9), Inches(height))
            box.text_frame.word_wrap = True
            box.text_frame.text = content
            for paragraph in box.text_frame.paragraphs:
                paragraph.font.name = "Microsoft YaHei"
                paragraph.font.size = Pt(size)
                paragraph.font.bold = bold
    presentation.core_properties.title = packet["title"]
    output = BytesIO()
    presentation.save(output)
    reopened = Presentation(BytesIO(output.getvalue()))
    if len(reopened.slides) != len(packet["pages"]):
        raise ToolError("EXPORT_FAILED", "PPT 成品读回校验失败")
    return output.getvalue()


def _render_pdf(packet: dict) -> bytes:
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.pdfgen import canvas
    import pypdfium2 as pdfium

    if not FONT_PATH.is_file():
        raise ToolError("FONT_UNAVAILABLE", "离线中文字体缺失，无法生成 PDF")
    if "OrviaNotoSC" not in pdfmetrics.getRegisteredFontNames():
        pdfmetrics.registerFont(TTFont("OrviaNotoSC", str(FONT_PATH)))
    output = BytesIO()
    pdf = canvas.Canvas(output, pagesize=A4, pageCompression=1)
    pdf.setTitle(packet["title"])

    def lines(value: str, size: int, width: float):
        current = ""
        for char in value:
            if char == "\n":
                yield current
                current = ""
            elif pdfmetrics.stringWidth(current + char, "OrviaNotoSC", size) > width and current:
                yield current
                current = char
            else:
                current += char
        if current:
            yield current

    for index, page in enumerate(packet["pages"]):
        pdf.setFont("OrviaNotoSC", 17)
        for title_line, line in enumerate(lines(packet["title"], 17, 495)):
            pdf.drawString(48, 801 - title_line * 21, line)
        pdf.setFont("OrviaNotoSC", 15)
        pdf.drawString(48, 745, page["heading"])
        pdf.setFont("OrviaNotoSC", 9)
        pdf.drawString(48, 720, page["label"])
        body = page["body"] or "\n".join(f"[{r['number']}] {r['title']} · {r['locator']}\n{r['citation']}" for r in page.get("references", []))
        y = 687
        pdf.setFont("OrviaNotoSC", 11)
        for line in lines(body, 11, 495):
            if y < 115:
                raise ToolError("LAYOUT_OVERFLOW", "PDF 内容超出单页预算，请缩短文字后重试")
            pdf.drawString(48, y, line)
            y -= 18
        if page["numbers"]:
            pdf.drawString(48, y - 14, "引用：" + " ".join(f"[{n}]" for n in page["numbers"]))
        pdf.setFont("OrviaNotoSC", 8)
        for row, line in enumerate(lines(packet["notice"], 8, 495)):
            pdf.drawString(48, 81 - row * 12, line)
        pdf.drawRightString(545, 48, f"{index + 1} / {len(packet['pages'])}")
        pdf.showPage()
    pdf.save()
    result = output.getvalue()
    with closing(pdfium.PdfDocument(result)) as reopened:
        if len(reopened) != len(packet["pages"]):
            raise ToolError("EXPORT_FAILED", "PDF 成品读回校验失败")
    return result


def render_publication(packet: dict) -> bytes:
    """固定格式分发；不接受模板、脚本、URL 或本地素材路径。"""
    output = {"docx": _render_docx, "pptx": _render_pptx, "pdf": _render_pdf}[packet["format"]](packet)
    if len(output) > 2 * 1024 * 1024:
        raise ToolError("OUTPUT_LIMIT", "成品超过 2 MiB 写入预算")
    return output
