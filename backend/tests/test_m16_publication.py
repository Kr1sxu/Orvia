"""M16 合成简报：真实本地渲染/SQLite/文件网关；Main 回答由固定 mock 产生。"""

import asyncio
from contextlib import closing
from io import BytesIO
import json
from uuid import uuid4

from docx import Document
import pypdfium2 as pdfium
from pptx import Presentation

from orvia_backend.configuration.client import Completion
from test_chat import call, create, setup
from test_m15_synthesis import seed


def test_three_publications_readback_version_save_isolation(tmp_path):
    async def run():
        app = await setup(tmp_path)
        try:
            cid = (await create(app))["id"]
            other = (await create(app))["id"]
            sources = await seed(app, cid)
            synthesis = {"id": cid, "mode": "summary", "question": "概括合成资料", "sources": sources[:1]}
            evidence = (await call(app, "chat.synthesis.preview", synthesis))["result"]
            cite = evidence["fragments"][0]["citation"]
            calls = []

            async def complete(profile, messages, **kwargs):
                calls.append(profile.model)
                return Completion(json.dumps({"answer": "合成摘要：需要保留来源和人工核对。",
                    "claims": [{"text": "合成资料要求保留来源。", "kind": "fact", "citations": [cite]}]}, ensure_ascii=False), (), "stop", {"total_tokens": 42})

            app.chat.client.complete = complete
            generated = (await call(app, "chat.synthesis.generate", {**synthesis, "revision": evidence["revision"], "request_id": str(uuid4())}))["result"]
            message = generated["messages"][-1]
            assert message["kind"] == "synthesis" and calls == ["deepseek-flash"]
            for format in ("docx", "pptx", "pdf"):
                request = {"id": cid, "message_id": message["id"], "format": format, "title": "合成资料简报",
                           "answer": "已确认的合成说明。", "claim_texts": ["合成资料需要核对，保留来源。"]}
                denied = await call(app, "chat.publication.preview", {**request, "id": other})
                assert denied["error"]["code"] == "NOT_FOUND"
                assert not (await call(app, "chat.publication.preview", {**request, "path": str(tmp_path / "bad")}))["ok"]
                preview = (await call(app, "chat.publication.preview", request))["result"]
                assert preview["source_revision"] and preview["references"][0]["citation"] == cite
                assert "合成资料需要核对" in "".join(page["body"] for page in preview["pages"])
                destination = tmp_path / f"brief.{format}"
                save = {**request, "revision": preview["revision"], "request_id": str(uuid4()), "path": str(destination)}
                stale = await call(app, "chat.publication.save", {**save, "revision": "0" * 64})
                assert stale["error"]["code"] == "STALE_APPROVAL" and not destination.exists()
                finished = (await call(app, "chat.publication.save", save))["result"]
                assert finished["messages"][-1]["kind"] == "publication"
                assert finished["messages"][-1]["data"]["filename"] == destination.name
                data = destination.read_bytes()
                if format == "docx":
                    document = Document(BytesIO(data))
                    assert "合成资料简报" in "\n".join(par.text for par in document.paragraphs)
                    assert document.tables and cite in document.tables[0].rows[1].cells[1].text
                elif format == "pptx":
                    slides = Presentation(BytesIO(data)).slides
                    assert len(slides) == len(preview["pages"])
                    assert "合成资料需要核对" in "\n".join(shape.text for slide in slides for shape in slide.shapes if shape.has_text_frame)
                else:
                    with closing(pdfium.PdfDocument(data)) as pdf:
                        assert len(pdf) == len(preview["pages"])
                        with closing(pdf[0]) as page, closing(page.get_textpage()) as text:
                            assert "合成资料简报" in text.get_text_range()
                duplicate = (await call(app, "chat.publication.save", save))["result"]
                assert duplicate == finished and destination.read_bytes() == data
                changed_target = await call(app, "chat.publication.save", {**save, "path": str(tmp_path / f"other.{format}")})
                assert changed_target["error"]["code"] == "CONFLICT" and not (tmp_path / f"other.{format}").exists()
                rejected = (await call(app, "chat.publication.save", {**save, "request_id": str(uuid4())}))["result"]
                assert rejected["messages"][-1]["data"]["code"] == "EXPORT_EXISTS"
                assert destination.read_bytes() == data
                wrong = tmp_path / f"wrong.{format}.txt"
                wrong_result = (await call(app, "chat.publication.save", {**save, "request_id": str(uuid4()), "path": str(wrong)}))["result"]
                assert wrong_result["messages"][-1]["data"]["code"] == "EXPORT_INVALID" and not wrong.exists()
            assert calls == ["deepseek-flash"]
        finally:
            await app.close()
    asyncio.run(run())


def test_long_text_pages_and_invalid_contents(tmp_path):
    async def run():
        app = await setup(tmp_path)
        try:
            cid = (await create(app))["id"]
            await app.chat.repository.append(cid, "assistant", "synthetic", "synthesis", {
                "answer": "摘要" * 1000,
                "claims": [{"text": "结论" * 250, "kind": "unknown", "citations": []}],
                "citations": [], "coverage": [], "revision": "a" * 64, "model": "deepseek-flash", "usage": {}})
            message = (await call(app, "chat.get", {"id": cid}))["result"]["messages"][-1]
            req = {"id": cid, "message_id": message["id"], "format": "pptx", "title": "长文测试",
                   "answer": "摘要" * 1000, "claim_texts": ["结论" * 250]}
            preview = (await call(app, "chat.publication.preview", req))["result"]
            assert len(preview["pages"]) > 10
            assert "".join(p["body"] for p in preview["pages"]).count("摘要") >= 1000
            invalid = await call(app, "chat.publication.preview", {**req, "title": "含\x01控制"})
            assert invalid["error"]["code"] == "INVALID_CONTENT"
        finally:
            await app.close()
    asyncio.run(run())
