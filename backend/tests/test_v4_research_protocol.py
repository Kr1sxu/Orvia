"""V4-010真实协议、HTTP/SQLite/原文引用/三格式成品；云模型仅合成替身。"""
import asyncio
from contextlib import closing
from io import BytesIO
import json
import io
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
import threading
from uuid import uuid4

from docx import Document
from pptx import Presentation
import pypdfium2 as pdfium
import pytest

from test_chat import setup, create, call
from orvia_backend.browser.network import SafeHTTP
from orvia_backend.context import ContextService
from v4_research_fixture import LocalSocketTransport, dns
from orvia_backend.configuration.client import Completion


async def result(app, method, params):
    response = await call(app, method, params)
    assert response['ok'], response
    return response['result']


def test_private_contract_rejects_authority_bounds_and_cross_conversation(tmp_path):
    async def run():
        app = await setup(tmp_path)
        try:
            cid = (await create(app))['id']
            body = {'id': cid, 'question': '合成对比', 'urls': ['https://example.com/']}
            for extra in ({'approved': True}, {'urls': ['https://example.com/'] * 11}, {'queries': ['合成'] * 3}, {'sites': ['example.com'] * 6}, {'model': 'other'}, {'question': True}, {'urls': ['http://127.0.0.1/']}, {'sources': [{'kind':'browser','evidence_id':'bad'}]}):
                assert not (await call(app, 'research.create', {**body, **extra}))['ok']
            task = await result(app, 'research.create', body)
            cid2 = (await create(app))['id']
            assert not (await call(app, 'research.status', {'id': cid2, 'operation_id':task['operation_id']}))['ok']
            assert not (await call(app, 'research.generate', {'id':cid,'operation_id':task['operation_id'],'stage':'final','revision':'a'*64,'approved':True}))['ok']
            assert task['state'] == 'planned' and task['pages'] == []
        finally:
            await app.close()
    asyncio.run(run())


def test_actual_http_citations_publication_three_formats_restart_and_purge(tmp_path):
    calls = []
    class Handler(BaseHTTPRequestHandler):
        def do_GET(self):
            calls.append(self.path)
            content = ('合成实验对比：记录A为10，记录B为20。只有本次合成观察，缺少长期测量。' + self.path).encode()
            self.send_response(200); self.send_header('Content-Type', 'text/plain; charset=utf-8')
            self.send_header('Content-Length', str(len(content))); self.end_headers(); self.wfile.write(content)
        def log_message(self, *_args):
            pass
    server = ThreadingHTTPServer(('127.0.0.1',0), Handler)
    thread = threading.Thread(target=server.serve_forever, daemon=True); thread.start()
    async def run():
        app = await setup(tmp_path)
        try:
            cid = (await create(app))['id']
            app.browser.network = SafeHTTP(transport=LocalSocketTransport(server.server_port), resolver=dns)
            task = await result(app, 'research.create', {'id':cid,'question':'合成实验比较','urls':['http://research.example/a','http://research.example/a','http://research.example/b'],'queries':['合成缺口'],'sites':['research.example']})
            request = {'id':cid,'operation_id':task['operation_id']}
            assert calls == []
            collected = await result(app, 'research.collect', request)
            assert calls == ['/a','/b']
            assert collected['state'] == 'limited' and collected['coverage']['search_unavailable']
            assert collected['coverage']['attempted_pages'] == 2
            assert not (await call(app, 'research.collect', request))['ok']
            packet = await result(app,'research.preview',{**request,'stage':'final'})
            assert len(packet['source_ids']) == 2 and all('合成' in f['text'] for f in packet['fragments'])
            assert packet['bytes'] == len((packet['system']+packet['input']).encode())
            model_calls = []
            async def complete(profile,messages,**kwargs):
                assert profile.model == 'deepseek-flash' and profile.base_url == 'https://api.deepseek.com'
                assert messages == [{'role':'system','content':packet['system']},{'role':'user','content':packet['input']}]
                model_calls.append(profile.model)
                fragments = packet['fragments']
                return Completion(json.dumps({'answer':'比较：两份合成资料。冲突：记录不同。缺口：缺少长期測量。覆盖：所选原文片段。','claims':[{'text':'两份记录的数值相互冲突。','kind':'conflict','citations':[f['citation'] for f in fragments[:2]]},{'text':'长期效果无法确认。','kind':'unknown','citations':[]}]}),(), 'stop', {})
            app.chat.client.complete = complete
            finished = await result(app,'research.generate',{**request,'stage':'final','revision':packet['revision']})
            assert finished['final']['state'] == 'saved' and model_calls == ['deepseek-flash']
            message = await app.chat.repository.message(cid, finished['final']['message_id']); data=message['data']
            for citation in data['citations']:
                source = await app.chat.evidence.get(cid,citation['evidence_id'])
                assert citation['locator'] == source['source_url'] and source['content_hash']
            for format in ('docx','pptx','pdf'):
                req = {'id':cid,'message_id':message['id'],'format':format,'title':'合成调研简报','answer':data['answer'],'claim_texts':[c['text'] for c in data['claims']]}
                preview = await result(app,'chat.publication.preview',req)
                path = tmp_path / ('合成调研.' + format)
                snapshot = await result(app,'chat.publication.save',{**req,'revision':preview['revision'],'request_id':str(uuid4()),'path':str(path)})
                assert snapshot['messages'][-1]['kind'] == 'publication'
                content=path.read_bytes()
                if format=='docx':
                    doc=Document(BytesIO(content)); assert doc.tables and data['citations'][0]['citation'] in doc.tables[0].rows[1].cells[1].text
                elif format=='pptx':
                    presentation=Presentation(BytesIO(content)); assert len(presentation.slides)==len(preview['pages'])
                    assert any('合成调研简报' in shape.text for slide in presentation.slides for shape in slide.shapes if shape.has_text_frame)
                else:
                    with closing(pdfium.PdfDocument(content)) as pdf:
                        assert len(pdf)==len(preview['pages'])
                        with closing(pdf[0]) as page, closing(page.get_textpage()) as text:
                            assert '合成调研简报' in text.get_text_range()
            task=await result(app,'research.status',request)
            assert {item['format'] for item in task['publications']}=={'docx','pptx','pdf'}
            assert all(item['verified'] for item in task['publications'])
            await app.close(); app=await setup(tmp_path)
            restored = await result(app,'research.status',request)
            assert restored==task and calls==['/a','/b']
            assert not (await call(app,'research.generate',{**request,'stage':'final','revision':packet['revision']}))['ok']
            await result(app,'chat.delete',{'id':cid})
            for table in ('research_tasks','research_sources','research_attempts'):
                async with app.store._db().execute('SELECT count(*) FROM '+table+' WHERE cid=?',(cid,)) as cursor:
                    assert (await cursor.fetchone())[0]==0
            assert all((tmp_path/('合成调研.'+format)).exists() for format in ('docx','pptx','pdf'))
            with pytest.raises(ValueError, match='已删除'):
                await ContextService(app.store).index_text(cid, 'browser:late', '合成晚到原文')
        finally:
            await app.close()
    try:
        asyncio.run(run())
    finally:
        server.shutdown();server.server_close();thread.join(2)


def test_standalone_mission_context_still_supported(tmp_path):
    async def run():
        app=await setup(tmp_path)
        try:
            mission=await result(app,'missions.create',{'client_request_id':str(uuid4()),'title':'合成独立目录任务'})
            context=ContextService(app.store)
            await context.index_text(mission['id'],'synthetic','合成原文检索')
            assert (await context.search(mission['id'],'合成'))['evidence']
        finally:await app.close()
    asyncio.run(run())


def test_publication_callback_uses_exact_receipt_despite_later_message(tmp_path):
    """保存后并发消息不能使调研关联误取最近消息；真实DOCX写入，零模型。"""
    async def run():
        app=await setup(tmp_path)
        try:
            cid=(await create(app))['id']
            await app.chat.repository.append(cid,'assistant','合成已保存结果','synthesis',{'answer':'合成缺口待补充','claims':[{'text':'暂无充分依据','kind':'unknown','citations':[]}],'citations':[],'coverage':[],'revision':'a'*64,'model':'deepseek-flash','usage':{}})
            messages,_=await app.chat.repository.messages(cid);message=messages[-1]
            append=app.chat.repository.append;received=[]
            async def interleaved(conversation_id,role,text,kind='text',data=None):
                await append(conversation_id,role,text,kind,data)
                if kind=='publication':await append(conversation_id,'system','合成同时状态消息','text')
            async def capture(conversation_id,message_id,data):
                assert conversation_id==cid and message_id==message['id']
                received.append(data)
            app.chat.repository.append=interleaved;app.research.record_publication=capture
            req={'id':cid,'message_id':message['id'],'format':'docx','title':'合成并发简报','answer':message['data']['answer'],'claim_texts':['暂无充分依据']}
            packet=await result(app,'chat.publication.preview',req);path=tmp_path/'合成并发.docx';rid=str(uuid4())
            snapshot=await result(app,'chat.publication.save',{**req,'revision':packet['revision'],'path':str(path),'request_id':rid})
            assert snapshot['messages'][-1]['kind']=='text' and path.exists()
            assert received[0]['filename']==path.name and received[0]['format']=='docx' and received[0]['request_id']==rid
            saved=next(m for m in snapshot['messages'] if m['kind']=='publication')
            assert received[0]==saved['data']
        finally:await app.close()
    asyncio.run(run())


def test_actual_http_inflight_cancel_blocks_delete_then_rejects_late_body(tmp_path):
    started=threading.Event(); released=threading.Event(); requests=[]
    class Handler(BaseHTTPRequestHandler):
        def do_GET(self):
            requests.append(self.path);started.set();released.wait(5)
            try:
                self.send_response(200);self.send_header('Content-Type','text/plain; charset=utf-8');self.end_headers();self.wfile.write('合成晚到结果'.encode())
            except (BrokenPipeError,ConnectionResetError):
                pass
        def log_message(self,*_args):pass
    server=ThreadingHTTPServer(('127.0.0.1',0),Handler);thread=threading.Thread(target=server.serve_forever,daemon=True);thread.start()
    async def run():
        app=await setup(tmp_path)
        try:
            cid=(await create(app))['id'];app.browser.network=SafeHTTP(transport=LocalSocketTransport(server.server_port),resolver=dns)
            task=await result(app,'research.create',{'id':cid,'question':'合成取消','urls':['http://research.example/slow','http://research.example/next']})
            request={'id':cid,'operation_id':task['operation_id']};pending=asyncio.create_task(call(app,'research.collect',request))
            for _ in range(200):
                if started.is_set():break
                await asyncio.sleep(.01)
            assert started.is_set()
            assert (await call(app,'chat.delete',{'id':cid}))['error']['code']=='DELETE_BLOCKED'
            cancelled=await result(app,'research.cancel',request)
            assert cancelled['state']=='cancelled' and (await pending)['ok']
            released.set();await asyncio.sleep(.05)
            assert requests==['/slow']
            async with app.store._db().execute('SELECT count(*) FROM browser_evidence WHERE mission_id=?',(cid,)) as cursor:
                assert (await cursor.fetchone())[0]==0
            assert not (await call(app,'research.collect',request))['ok']
            await result(app,'chat.delete',{'id':cid})
        finally:released.set();await app.close()
    try:asyncio.run(run())
    finally:released.set();server.shutdown();server.server_close();thread.join(2)


@pytest.mark.parametrize('control',['research.status','research.history','research.cancel'])
def test_private_stdio_research_control_available_during_collect(monkeypatch,control):
    """实际调度及输出：取消/只读状态旁路，新的写步骤仍按原顺序执行。"""
    from orvia_backend import server
    class Controlled:
        def __init__(self,event_sink):self.started=asyncio.Event();self.observed=asyncio.Event()
        async def handle(self,line):
            request=json.loads(line);method=request['method']
            if method=='research.collect':self.started.set();await asyncio.wait_for(self.observed.wait(),1)
            elif method==control:await asyncio.wait_for(self.started.wait(),1);self.observed.set()
            elif method=='research.create':assert self.observed.is_set()
            return {'v':1,'id':request['id'],'ok':True,'result':{'method':method}}
        async def close(self):pass
    monkeypatch.setattr(server,'Application',Controlled)
    reader=io.BytesIO(b''.join((json.dumps({'v':1,'id':str(uuid4()),'method':method,'params':{}})+'\n').encode() for method in ['research.collect','research.create',control]));writer=io.BytesIO()
    asyncio.run(server.serve(reader,writer));responses=[json.loads(line) for line in writer.getvalue().splitlines()]
    assert len(responses)==3 and all(r['ok'] for r in responses)
    order=[r['result']['method'] for r in responses];assert order.index(control)<order.index('research.create')
