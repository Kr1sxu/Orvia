"""独立检查M20实际保存的三格式成品；只报告结构、引用计数和文件hash。"""

import argparse
from contextlib import closing
import hashlib
import json
from pathlib import Path


def compact(text):
    return "".join(text.split())


def inspect(directory, expected):
    from docx import Document
    from pptx import Presentation
    import pypdfium2 as pdfium

    facts = []
    for format in ("docx", "pptx", "pdf"):
        source = directory / f"real-brief.{format}"
        data = source.read_bytes()
        assert 1000 < len(data) <= 2 * 1024 * 1024
        if format == "docx":
            document = Document(source)
            text = "\n".join([*(p.text for p in document.paragraphs),
                              *(cell.text for table in document.tables for row in table.rows for cell in row.cells)])
            pages = sum("lastRenderedPageBreak" in p._p.xml or "w:type=\"page\"" in p._p.xml for p in document.paragraphs) + 1
        elif format == "pptx":
            presentation = Presentation(source)
            pages = len(presentation.slides)
            text = "\n".join(shape.text for slide in presentation.slides for shape in slide.shapes if shape.has_text_frame)
        else:
            with closing(pdfium.PdfDocument(source)) as pdf:
                pages = len(pdf)
                parts = []
                for index in range(pages):
                    with closing(pdf[index]) as page:
                        with closing(page.get_textpage()) as layer:
                            parts.append(layer.get_text_range())
                        # 导出逐页合成成品图，只供本轮视觉复核，不捕获系统桌面。
                        with closing(page.render(scale=1.3)) as bitmap:
                            bitmap.to_pil().save(directory / f"real-brief-pdf-{index + 1}.png")
                text = "\n".join(parts)
        visible = compact(text)
        assert "资料简报" in visible and "来源与引用" in visible
        assert 1 <= pages <= 48
        assert expected and all(compact(cite) in visible for cite in expected)
        facts.append({"format": format, "bytes": len(data), "sha256": hashlib.sha256(data).hexdigest(),
                      "pagesOrPlannedBreaks": pages, "titlePresent": True,
                      "citationCount": len(expected), "citationsPresent": True, "readBackVerified": True})
    return facts


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--directory", required=True)
    parser.add_argument("--expected", required=True)
    arguments = parser.parse_args()
    # 审计只接受本仓库M20产物；不是产品运行时或安装版解释器回退入口。
    results = Path(__file__).resolve().parents[2] / "artifacts/test-results/M20"
    directory = Path(arguments.directory).resolve(strict=True)
    expected = Path(arguments.expected).resolve(strict=True)
    assert directory.is_relative_to(results) and expected.parent == directory
    print(json.dumps(inspect(directory, json.loads(expected.read_text(encoding="utf-8"))), ensure_ascii=False))
