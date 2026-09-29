"""显式执行一次真实 Main 合成调用，再用其已保存回答本地生成三格式简报。"""

import argparse
import asyncio
from contextlib import closing
from io import BytesIO
import json
from pathlib import Path
import sys
from tempfile import TemporaryDirectory
from uuid import uuid4

from docx import Document
import pypdfium2 as pdfium
from pptx import Presentation

from live_m15_synthesis import call, development_key
from orvia_backend.application import Application


ROOT = Path(__file__).resolve().parents[2]
RESULTS = ROOT / "artifacts/test-results/M16"


async def run() -> dict:
    """只保留计数和状态；合成资料、模型原文和 Key 仅在临时进程/目录内。"""
    key = development_key()
    if not key:
        return {"status": "MISSING_CREDENTIAL", "calls": 0}
    RESULTS.mkdir(parents=True, exist_ok=True)
    with TemporaryDirectory(prefix="live-integration-", dir=RESULTS) as directory:
        app = Application()
        try:
            await call(app, "hello")
            initialized = await call(app, "initialize", {"data_directory": directory, "credentials": {"main": key}})
            if not initialized.get("ok"):
                return {"status": "INIT_FAILED", "calls": 0}
            cid = (await call(app, "chat.create", {"client_request_id": str(uuid4()), "title": "M16 synthetic integration"}))["result"]["id"]
            mission = await app.store.get_mission(cid)
            main = next(profile for profile in mission.models if profile.role == "main")
            if (main.model, main.base_url) != ("deepseek-flash", "https://api.deepseek.com"):
                return {"status": "PROFILE_MISMATCH", "calls": 0}

            # 文档/网页均为本进程构造的短合成证据；不访问用户目录或真实网页。
            parsed = {"format": "docx", "units": [{"number": 1, "locator": "段落 1",
                "text": "合成项目乙的交付日期是2032年6月12日，验收备注为保留来源。",
                "method": "text", "confidence": None, "error": None}],
                "total_units": 1, "truncated": False, "missing_units": [], "error": None}
            document = await app.chat.documents.save(cid, "synthetic-m16.docx", b"m16-synthetic-only", parsed)
            browser = await app.chat.evidence.save(cid, {"title": "合成网页", "source_url": "https://example.com/m16-synthetic",
                "mode": "http", "accessed_at": "2026-09-29", "content": "合成网页也记录项目乙在2032年6月12日交付。",
                "truncated": False, "error": None})
            synthesis = {"id": cid, "mode": "answer", "question": "综合两个合成来源说明项目乙交付日期并标注来源。",
                "sources": [{"kind": "document", "evidence_id": document["evidence_id"]},
                            {"kind": "browser", "evidence_id": browser["evidence_id"]}]}
            preview = (await call(app, "chat.synthesis.preview", synthesis))["result"]
            if len(preview["fragments"]) < 2:
                return {"status": "INSUFFICIENT_FRAGMENTS", "calls": 0}

            # 生产适配器固定 20 秒 HTTP 超时、30 秒总等待、1024 输出 token、零自动重试。
            generated = await call(app, "chat.synthesis.generate", {**synthesis, "revision": preview["revision"],
                                                                     "request_id": str(uuid4())})
            if not generated.get("ok"):
                return {"status": "GENERATE_PROTOCOL_FAILED", "calls": 1}
            message = generated["result"]["messages"][-1]
            if message["kind"] != "synthesis":
                return {"status": message.get("data", {}).get("code", "MODEL_RESULT_FAILED"), "calls": 1}
            data = message["data"]
            if not data["claims"] or not data["citations"]:
                return {"status": "CITATIONS_MISSING", "calls": 1}

            formats = []
            for format in ("docx", "pptx", "pdf"):
                request = {"id": cid, "message_id": message["id"], "format": format,
                           "title": "合成项目乙简报", "answer": data["answer"],
                           "claim_texts": [claim["text"] for claim in data["claims"]]}
                planned = await call(app, "chat.publication.preview", request)
                if not planned.get("ok"):
                    return {"status": "PUBLICATION_PREVIEW_FAILED", "calls": 1, "format": format}
                packet = planned["result"]
                target = Path(directory) / f"live-m16.{format}"
                saved = await call(app, "chat.publication.save", {**request, "revision": packet["revision"],
                                                                    "request_id": str(uuid4()), "path": str(target)})
                if not saved.get("ok") or saved["result"]["messages"][-1]["kind"] != "publication":
                    return {"status": "PUBLICATION_SAVE_FAILED", "calls": 1, "format": format}
                content = target.read_bytes()
                if format == "docx":
                    document = Document(BytesIO(content))
                    opened = bool(document.tables) and packet["title"] in "\n".join(p.text for p in document.paragraphs)
                elif format == "pptx":
                    slides = Presentation(BytesIO(content)).slides
                    opened = len(slides) == len(packet["pages"])
                else:
                    with closing(pdfium.PdfDocument(content)) as pdf:
                        opened = len(pdf) == len(packet["pages"])
                if not opened:
                    return {"status": "READBACK_FAILED", "calls": 1, "format": format}
                formats.append(format)
            return {"status": "PASS", "calls": 1, "model": main.model, "sources": 2,
                    "claims": len(data["claims"]), "citations": len(data["citations"]),
                    "total_tokens": data["usage"].get("total_tokens", "unavailable"), "formats": formats}
        finally:
            await app.close()


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--run-live", action="store_true")
    args = parser.parse_args()
    if not args.run_live:
        print("M16 live: skipped (requires --run-live)")
        sys.exit(0)
    try:
        result = asyncio.run(run())
    except Exception as error:
        # 异常正文可能包含供应商返回内容，控制台只输出类型。
        result = {"status": "EXCEPTION", "type": type(error).__name__}
    RESULTS.mkdir(parents=True, exist_ok=True)
    (RESULTS / "live-integration.json").write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
    print("M16 live: " + json.dumps(result, ensure_ascii=False))
    sys.exit(0 if result["status"] == "PASS" else 1)
